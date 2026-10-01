#!/usr/bin/env python3
"""Build SHA256_Literature_Tracker.xlsx and papers.csv from the table below.

Edit PAPERS, then run:  python3 literature/build_tracker.py
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))

COLUMNS = [
    "ID", "Year", "Title", "Authors", "Venue", "Type",
    "Impact factor (approx.)", "Venue weight",
    "Problem", "What they did", "Main achievement",
    "Gap we find", "How our work affects them",
    "Relevance (1-5)", "Novelty threat (1-5)",
    "PDF", "Link", "Check status",
]

# Impact factors are approximate JCR values for journals. Conferences, preprints,
# patents and datasheets have no impact factor; the CORE rank or a plain tier is
# given instead. Check every IF on JCR before quoting it.
#
# Relevance: 5 = direct prior art for B/C/D or the operand-scheduling finding,
#            1 = background only.
# Threat:    5 = kills a novelty claim outright, 1 = no overlap.
PAPERS = [
    # ---- Recent unrolling / round-optimisation cluster (2022-2026) ----
    ("P01", 2022, "FPGA-based Implementation of SHA-256 with Improvement of Throughput using Unfolding Transformation",
     "S. Suhaili, N. Julai", "Pertanika J. Sci. & Technol. 30(1):581-603", "Journal",
     "No JCR IF (Scopus-indexed)", "Low-Med",
     "SHA-256 throughput; cycles per block",
     "Unfolding x2 / x4 on Arria II GX; 3 separately built designs",
     "x2: 64->34 cycles, Fmax ROSE; x4: 4196 Mbps (+58% vs iterative)",
     "Separate builds, no matched constraints, operand order not reported",
     "We re-test x2 under controlled build; operand order (0.500 vs 0.680) explains their Fmax rise",
     5, 3, "Open (Pertanika site)",
     "http://www.pertanika.upm.edu.my/pjst/browse/regular-issue?article=JST-2895-2021", "Verified (survey [6])"),
    ("P02", 2025, "Simulation-based power estimation for high throughput SHA-256 design on unfolding transformation",
     "S. Suhaili, N. Julai, A. Lit, M. Huja Husin, M.F. Mohd Sabri", "AIP Conf. Proc. 3056, 060007 (SIE2022)", "Conference",
     "None (proceedings)", "Low",
     "Power of unfolded SHA-256",
     "Unfolding + Gray-encoded control; Quartus/ModelSim power sim",
     "x4 with Gray encoding: dynamic power -43.4%",
     "Simulated power only; still no Fmax mechanism, still Altera",
     "Same family as P01; our operand-order mechanism applies to their x2/x4 too",
     3, 1, "Paywalled", "https://doi.org/10.1063/5.0209048", "Search-verified"),
    ("P03", 2023, "An Unrolled and Pipelined Architecture for SHA2 Family of Hash Functions on FPGA",
     "O.B. Gamgam", "ISCTurkiye 2023 (IEEE)", "Conference",
     "None (conf.)", "Low",
     "Throughput/slice of SHA-2",
     "Unrolling + Round Pipelined Technique (RPT); SHA-512 80->42 steps",
     "Throughput/slice +57% (SHA-256), +17% (SHA-512)",
     "One operating point; unroll and pipeline never varied independently",
     "Kills novelty of Config D; we add the controlled 2x2 comparison",
     5, 4, "Paywalled", "https://ieeexplore.ieee.org/document/10336096", "Verified (survey [7])"),
    ("P04", 2025, "An optimized hardware implementation of SHA-256 round computation",
     "M. Yao, Z. Xue, H. Li, S. Shen", "The Computer Journal 68(4):355-359", "Journal",
     "~1.5", "Med",
     "Round critical path",
     "Rearranged round; critical path split into 2 add stages; 4-2 compressor",
     "366 MHz, 1990 Mbps, 1.86 Mbps/slice",
     "Single iterative round only; no unrolled-pair analysis",
     "Confirms operand rearrangement is KNOWN -> we claim its measured effect on unrolling, not the technique",
     5, 4, "Paywalled", "https://academic.oup.com/comjnl/article-abstract/68/4/355/7899894", "Search-verified (NEW)"),
    ("P05", 2025, "Efficient FPGA-Based Implementation of SHA-256 Hash Function for Higher Throughput",
     "P.R. Lawhale, S.N. Kale", "ICSM 2024, LNEE vol. 1381 (Springer)", "Conference",
     "None (proceedings)", "Low",
     "SHA-256 throughput",
     "VHDL on Stratix-III, Quartus-II",
     "1090.512 Mbps",
     "No architectural variable studied",
     "Benchmark row only",
     2, 1, "Paywalled", "https://doi.org/10.1007/978-981-96-3644-0_26", "Verified (survey [8])"),
    ("P06", 2025, "A Constructive High-Speed Crypto-mining Approach with Dual SHA-256 on an FPGA",
     "V. Pavan Kumar, A. Alagarsamy", "VLSID 2025 (IEEE)", "Conference",
     "None (conf.)", "Med",
     "Mining hash speed/power",
     "Configurable dual SHA-256 datapath, Virtex-7",
     "+6.09% speed/power, +13.40% functional efficiency",
     "Single-digit gains; no Fmax mechanism",
     "Shows scale of credible 2025 gains; our mechanism is bigger (0.500 vs 0.680)",
     3, 2, "Paywalled", "https://ieeexplore.ieee.org/document/10900629", "Search-verified (authors now confirmed)"),
    ("P07", 2024, "SHA-256 Hardware Proposal for IoT Devices in the Blockchain Context",
     "Santos Jr., da Silva, Torquato, Silva, Fernandes", "Sensors 24(12):3908 (MDPI)", "Journal",
     "~3.4", "Med",
     "Low-power hashing for IoT blockchain",
     "1-16 clustered SHA-256 cores, Virtex-6",
     "~1.4 Gbps @16 cores, 0.503 W; Fmax ~11 MHz at 71% LUTs",
     "Whole-core replication collapses Fmax; no interleaving",
     "Supports our structural argument; C gives 2x without replicating cores",
     4, 2, "Open (MDPI)", "https://www.mdpi.com/1424-8220/24/12/3908", "Verified (survey [3])"),
    ("P08", 2025, "Implementation of SHA-256 Used in Bitcoin Mining on FPGA",
     "H. Shah, Agrawal, Shah", "SmartCom 2025, LNNS vol. 1463 (Springer)", "Conference",
     "None (proceedings)", "Low",
     "Mining SHA-256 on Zynq",
     "Modular pipelined VHDL on ZedBoard",
     "+27% throughput over software; 12.898 W on-chip (as reported)",
     "Same board, no unroll/interleave study, weak SW baseline",
     "Same board -> our controlled data is the stronger result on Zynq-7020",
     4, 2, "Paywalled", "https://doi.org/10.1007/978-981-96-7514-2_28", "Verified (survey [4])"),
    ("P09", 2026, "Comparative Analysis of HLS- and RTL-Based SHA-256 Accelerators on a Zynq FPGA",
     "Bal", "Int. J. Adv. Eng. Pure Sci. 38(2):272-280", "Journal",
     "No JCR IF", "Low",
     "HLS vs RTL productivity/QoR",
     "HLS and RTL SHA-256 as AXI IP on PYNQ-Z1 @100 MHz",
     "Both ~16x over SW; RTL more resource-efficient",
     "Fixed architecture; speedup depends on SW baseline",
     "Explains our ~3x vs their 16x (baseline); no overlap with unroll question",
     4, 1, "Open (DergiPark)", "https://dergipark.org.tr/en/pub/jeps/article/1864699", "Verified (survey [5])"),
    ("P10", 2025, "Hardware Design and Implementation of Secure Hash Algorithm Based on FPGA",
     "M.F. Al-Gailani", "J. Internet Serv. Inf. Secur. 15(2)", "Journal",
     "No JCR IF (Scopus)", "Low",
     "SHA-2 area/power",
     "SHA-256/512 on XC7VX330T, ISE 14.7",
     "SHA-256: 1.021 Gbps, 333 slice regs, 0.394 W",
     "Benchmark only",
     "Comparison-table row",
     2, 1, "Open (JISIS)", "https://jisis.org/wp-content/uploads/2025/07/2025.I2.015.pdf", "Search-verified"),
    ("P11", 2024, "Implementation of Efficient Low Power SHA-256 Algorithm",
     "K.S. Kulkarni, S. Rekha, S. Kumar K", "ICNEWS 2024 (IEEE)", "Conference",
     "None (conf.)", "Low",
     "Low-power SHA-256",
     "Low-power RTL SHA-256",
     "Power reduction (see paper)",
     "No throughput-architecture study",
     "Background only",
     2, 1, "Paywalled", "https://doi.org/10.1109/ICNEWS60873.2024.10730997", "Search-verified"),
    ("P12", 2024, "Custom ASIC Design for SHA-256 Using Open-Source Tools",
     "L.D. Franck, G.A. Ginja, J.P. Carmo, J.A. Afonso, M. Luppe", "Computers 13(1):9 (MDPI)", "Journal",
     "~2.6", "Med",
     "Open-source ASIC flow for SHA-256",
     "SHA-256 accelerator taped out with open EDA",
     "Working ASIC flow end-to-end",
     "ASIC, iterative only",
     "Background; our ngspice method is a lighter transistor-level check",
     2, 1, "Open (MDPI)", "https://www.mdpi.com/2073-431X/13/1/9", "Search-verified"),
    ("P13", 2023, "High Throughput and Energy Efficient SHA-2 ASIC Design for Continuous Integrity Checking Applications",
     "A. Koutra, V. Tenentes", "IEEE ETS 2023", "Conference",
     "None (conf.)", "Med",
     "Throughput + energy for always-on integrity checks",
     "SHA-2 ASIC design for continuous checking",
     "Higher throughput and energy efficiency (see paper)",
     "ASIC; no FPGA unroll/interleave analysis",
     "Background",
     2, 1, "Paywalled", "https://ieeexplore.ieee.org/document/10174095", "Search-verified"),
    ("P14", None, "A High-Throughput and Energy-Efficient SHA-256 Design using Approximate Arithmetic",
     "J. Baik, Y. Kim", "IEIE Trans. Smart Proc. & Computing", "Journal",
     "No JCR IF", "Low",
     "Energy of SHA-256 adders",
     "Split k-bit approximate adders",
     "Fewer resources, higher throughput (approximate)",
     "Approximate hash -> not FIPS 180-4 compliant",
     "Not comparable; we stay bit-exact (103+21 checks)",
     1, 1, "Open (IEIE)", "https://ieiespc.org/ieiespc/XmlViewer/f417617", "Search-verified; YEAR UNCONFIRMED"),
    ("P15", 2025, "LiteHash: Hash Functions for Resource-Constrained Hardware",
     "S.D. Achar, Thejaswini P, S. Nandi, S. Nandi", "ACM TECS 24(2), Art. 28", "Journal",
     "~2.0", "Med",
     "Cheaper hashing at the edge",
     "SHA-512-like approximate hash variants",
     "Area -16.86%, power -32.48%, up to 2 Gbps",
     "New (non-standard) hash, not SHA-256",
     "None",
     1, 1, "Paywalled", "https://doi.org/10.1145/3677181", "Search-verified"),
    ("P16", 2020, "A High-Throughput Hardware Implementation of SHA-256 Algorithm",
     "Chen, Li", "IEEE ISCAS 2020", "Conference",
     "None (CORE B)", "Med",
     "SHA-256 throughput",
     "High-throughput SHA-256 datapath",
     "See paper",
     "No controlled unroll study",
     "Comparison-table row",
     2, 1, "Paywalled", "https://ieeexplore.ieee.org/document/9181065", "Search-snippet only; AUTHORS UNCONFIRMED"),
    # ---- Multi-message / interleaving (Config C) ----
    ("P17", 2026, "SHARMONY: Composing SHA-2 and SHA-3 Hardware for Crypto-Agile PQC",
     "L. Anwar, C.A. Lara-Nino, J.-Y. Park, M. Hutter", "IACR TCHES 2026; ePrint 2026/1572", "Journal (IACR)",
     "No JCR IF (top crypto-HW venue)", "High",
     "One engine for SHA-2 + SHA-3 in PQC",
     "Unified datapath; DUET mode = 2 independent SHA-256 streams in 64-bit lanes",
     "1656 Mbps SHA-256; beats OpenTitan/Caliptra/SPHINCSLET/SLotH engines",
     "No unroll-vs-interleave Fmax study",
     "Another KILL for Config C novelty (2026, top venue); cite it",
     5, 5, "Open (ePrint)", "https://eprint.iacr.org/2026/1572", "Search-verified (NEW)"),
    ("P18", 2012, "Simultaneous Hashing of Multiple Messages",
     "S. Gueron, V. Krasnov", "J. Information Security 3(4):319-325; ePrint 2012/371", "Journal",
     "No JCR IF", "Med (highly cited)",
     "Hash many independent messages fast",
     "4-buffer SIMD SHA-256 (S-HASH)",
     "~5.2 cycles/byte, 3.42x over OpenSSL 1.0.1",
     "Software/SIMD, not FPGA",
     "Multi-message idea is old -> C not novel",
     4, 4, "Open (ePrint)", "https://eprint.iacr.org/2012/371", "Verified (patent doc)"),
    ("P19", 2023, "Parallel SHA-256 on SW26010 many-core processor for hashing of multiple messages",
     "Z. Wang, X. Dong, Y. Kang, H. Chen", "J. Supercomputing 79:2332-2355", "Journal",
     "~2.5", "Med",
     "Multi-message SHA-256 throughput",
     "Data/instruction-level parallel SHA-256 on Sunway",
     "5.87 cpb single core (8.18x OpenSSL); 60.21 GB/s",
     "Many-core software",
     "Multi-message parallelism standard across platforms",
     3, 2, "Paywalled", "https://doi.org/10.1007/s11227-022-04750-7", "Search-verified"),
    ("P20", 2021, "A High-Performance Multimem SHA-256 Accelerator for Society 5.0",
     "T.H. Tran, H.L. Pham, Y. Nakashima", "IEEE Access 9:39182-39192", "Journal",
     "~3.4", "Med",
     "Hash many messages with high throughput",
     "Multi-memory SHA-256 accelerator",
     "High throughput multi-message FPGA accelerator",
     "No controlled unroll comparison",
     "Prior art for multi-message hardware",
     3, 3, "Open (IEEE Access)", "https://ieeexplore.ieee.org/document/9367201", "Search-verified"),
    ("P21", 2022, "A High-Efficiency FPGA-Based Multimode SHA-2 Accelerator",
     "H.L. Pham, T.H. Tran, V.T.D. Le, Y. Nakashima", "IEEE Access (2022)", "Journal",
     "~3.4", "Med",
     "One accelerator for all SHA-2 modes",
     "Multimode SHA-2 FPGA accelerator",
     "High efficiency across SHA-2 variants",
     "No unroll/interleave isolation",
     "Background",
     3, 2, "Open (IEEE Access)", "https://www.researchgate.net/publication/358093691", "Search-verified; DOI TO ADD"),
    ("P22", 2020, "Double SHA-256 Hardware Architecture With Compact Message Expander for Bitcoin Mining",
     "H.L. Pham, T.H. Tran, T.D. Phan, L.V.T. Duong, D.K. Lam, Y. Nakashima", "IEEE Access 8:139634-139646", "Journal",
     "~3.4", "Med",
     "Mining throughput + power",
     "Fully unrolled double SHA-256 + compact message expander",
     "340 Gbps on ZCU102; 61.44 Gbps in 180 nm ASIC",
     "Fully unrolled+pipelined over independent nonces; no iterative-vs-unrolled Fmax study",
     "Background for unrolling at scale",
     3, 2, "Open (IEEE Access)", "https://doi.org/10.1109/ACCESS.2020.3012581", "Search-verified"),
    # ---- SLH-DSA / PQC consumers of SHA-256 ----
    ("P23", 2024, "Accelerating SLH-DSA by Two Orders of Magnitude with a Single Hash Unit (SLotH)",
     "M.-J.O. Saarinen", "CRYPTO 2024 (LNCS)", "Conference",
     "None (CORE A*)", "High",
     "SLH-DSA too slow on RoT SoCs",
     "One SHA2/SHAKE hash unit in a RoT",
     "~100x vs unaccelerated; 128f sign 4.9M cycles",
     "One hash unit; no multi-message core",
     "Our C/D directly fits SLH-DSA's independent-hash workload",
     4, 2, "Open (ePrint)", "https://eprint.iacr.org/2024/367", "Search-verified"),
    ("P24", 2025, "SPHINCSLET: An Area-Efficient Accelerator for the Full SPHINCS+ Digital Signature Algorithm",
     "Sahu et al. (confirm full list)", "ACM TECS (2025); ePrint 2025/621", "Journal",
     "~2.0", "Med-High",
     "Area-efficient full SLH-DSA",
     "Parameterisable SHA-2/SHAKE SLH-DSA on Artix-7",
     "SHA-2: 2-4x sign speedup in 6K-15K LUTs",
     "Hash core treated as fixed block",
     "Motivates faster SHA-256 core; our data is directly usable",
     4, 2, "Open (ePrint)", "https://eprint.iacr.org/2025/621", "Verified (survey [2]); AUTHORS TO CONFIRM"),
    ("P25", 2026, "Trident: Efficient FPGA Acceleration of XMSS Tree in SLH-DSA",
     "T. Bao, J. Ennis, K. Morozov, J. Xie", "IEEE FCCM 2026; ePrint 2026/837", "Conference",
     "None (top FPGA conf.)", "High",
     "WOTS+ keygen dominates SLH-DSA signing",
     "FPGA XMSS-tree accelerator, SHA-2 + SHAKE",
     "Efficient XMSS acceleration (see paper)",
     "Parallelism at tree level, not inside the SHA-256 core",
     "Our interleaved core could be its hash engine",
     3, 2, "Open (ePrint)", "https://eprint.iacr.org/2026/837", "Search-verified (NEW)"),
    ("P26", 2024, "A Configurable XMSS Post-Quantum Hardware Implementation with SM3 and SHA-256",
     "M. Ma, X. Ji, J. Cai", "ICCTIT 2024 (IEEE)", "Conference",
     "None (conf.)", "Low",
     "Configurable XMSS hardware",
     "XMSS with SM3/SHA-256 cores",
     "Configurable hash backend",
     "Hash core not studied",
     "Background",
     2, 1, "Paywalled", "https://doi.org/10.1109/ICCTIT64404.2024.10928459", "Search-verified"),
    ("P27", 2020, "FPGA-based Accelerator for Post-Quantum Signature Scheme SPHINCS-256",
     "D. Amiet, A. Curiger, P. Zbinden et al.", "IACR TCHES 2020", "Journal (IACR)",
     "No JCR IF (top crypto-HW venue)", "High",
     "Fast SPHINCS signing",
     "Kintex-7 SPHINCS-256 accelerator",
     "Sign 1.53 ms, verify 65 us",
     "Pre-standard SPHINCS; hash core not isolated",
     "Background",
     2, 1, "Open (TCHES)", "https://tches.iacr.org/index.php/TCHES/article/download/831/783/", "Search-verified; AUTHOR LIST TO CONFIRM"),
    # ---- SoC / RISC-V / Zynq integration ----
    ("P28", 2026, "Crypto-RV: High-Efficiency FPGA-Based RISC-V Cryptographic Co-Processor for IoT Security",
     "A.K. Pham, V.T. Vo, V.T.D. Le, T.H. Vu, H.L. Pham, V.T. Nguyen, Y. Nakashima", "arXiv 2602.04415", "Preprint",
     "None (preprint)", "Low-Med",
     "Many crypto primitives on one RISC-V",
     "Unified 64-bit datapath: SHA-256/512, SM3, SHA3, AES, Haraka",
     "ZCU102 @160 MHz; 165-1061x over base RISC-V",
     "Breadth, not depth on SHA-256 round",
     "Background",
     2, 1, "Open (arXiv)", "https://arxiv.org/abs/2602.04415", "Search-verified"),
    ("P29", 2024, "Assessing the Performance of OpenTitan as Cryptographic Accelerator in Secure Open-Hardware SoCs",
     "E. Parisi, A. Musa, M. Ciani, F. Barchi, D. Rossi, A. Bartolini, A. Acquaviva", "ACM CF'24", "Conference",
     "None (CORE B)", "Med",
     "Offload overhead of SoC crypto accelerators",
     "Benchmarks OpenTitan HMAC/SHA-256 offload",
     "HMAC core: 80 cycles/block; offload inefficiencies found",
     "Fixed engine; no architecture variation",
     "Same lesson as our 0.515 -> 0.624 system threshold",
     3, 1, "Open (arXiv)", "https://arxiv.org/abs/2402.10395", "Search-verified"),
    ("P30", 2021, "HW/SW Architecture Exploration for an Efficient Implementation of SHA-256",
     "M. Kammoun, M. Elleuchi, M. Abid, A.M. Obeid", "J. Commun. Softw. Syst. 17(2):87-96", "Journal",
     "Low (ESCI)", "Low",
     "LLS vs HLS SHA-256 on Zynq",
     "HW/SW co-design exploration on Zynq-7000",
     "Area/throughput/power trade-offs for LLS vs HLS",
     "No unroll-depth isolation",
     "Same platform class; background",
     3, 1, "Open (JCOMSS)", "https://jcoms.fesb.unist.hr/pdfs/v17n2_2021-0006_kammoun.pdf", "Search-verified"),
    # ---- Foundational (pre-2022) ----
    ("F01", 2006, "Optimisation of the SHA-2 Family of Hash Functions on FPGAs",
     "R. McEvoy, F. Crowe, C. Murphy, W. Marnane", "IEEE ISVLSI 2006, pp. 317-322", "Conference",
     "None (CORE B)", "Med",
     "SHA-2 throughput on FPGA",
     "2x-unrolled-pipelined SHA-2 core",
     "Unroll + pipeline combined in 2006",
     "No controlled axis isolation",
     "Kills Config D novelty; already our ref [11]",
     5, 5, "Paywalled", "https://ieeexplore.ieee.org/document/1602458", "Verified (survey [11])"),
    ("F02", 2006, "Improving SHA-2 Hardware Implementations",
     "R. Chaves, G. Kuzmanov, L. Sousa, S. Vassiliadis", "CHES 2006, LNCS 4249", "Conference",
     "None (CORE A)", "High",
     "Long T1 critical path",
     "Operation rescheduling across register boundary + CSA",
     ">50% better than commercial SHA-256 cores (Virtex-II Pro)",
     "Applied to iterative round, not inside an unrolled pair",
     "Technique behind our finding; we measure its effect on the unrolling verdict",
     5, 4, "Paywalled", "https://doi.org/10.1007/11894063_24", "Verified (survey [10])"),
    ("F03", 2004, "The design of a high speed ASIC unit for the hash function SHA-256 (384, 512)",
     "L. Dadda, M. Macchetti, J. Owen", "DATE 2004 Designers' Forum, pp. 70-76", "Conference",
     "None (CORE B)", "Med",
     "Adder delay in compressor/expander",
     "Carry-save adder trees",
     "Critical-path cut in 0.13 um",
     "ASIC; CSA not compared with operand order",
     "Alternative fix to the same adder chain",
     3, 2, "Paywalled", "https://www.researchgate.net/publication/4057195", "Verified (survey [12])"),
    ("F04", 2009, "A Top-Down Design Methodology for Ultrahigh-Performance Hashing Cores",
     "H. Michail, A. Kakarountas, A. Milidonis, C. Goutis", "IEEE TDSC 6(4):255-268", "Journal",
     "~7.0", "High",
     "Systematic unroll/pipeline of hash cores",
     "Top-down unroll + pipeline methodology",
     "Methodology for ultrahigh-throughput cores",
     "Comparisons span papers/devices",
     "Methodological ancestor; we tighten to one board/one constraint set",
     4, 3, "Paywalled", "https://doi.org/10.1109/TDSC.2008.15", "Verified (survey [13])"),
    ("F05", 2012, "On the Exploitation of a High-throughput SHA-256 FPGA Design for HMAC",
     "H. Michail et al. (confirm list)", "ACM TRETS 5(1)", "Journal",
     "~2.0", "Med",
     "High-throughput SHA-256 for HMAC",
     "Pre-computation / rescheduling in SHA-256 round",
     "High-throughput HMAC core",
     "Iterative round focus",
     "More prior art for operand pre-computation",
     3, 3, "Open (author repo)", "https://shura.shu.ac.uk/18348/1/harris.pdf", "Search-snippet; AUTHORS TO CONFIRM"),
    ("F06", 2019, "High-throughput and area-efficient fully-pipelined hashing cores using BRAM in FPGA",
     "Padhi, Chaudhari", "Microprocessors & Microsystems", "Journal",
     "~1.9", "Med",
     "Max throughput for SHA-1/256",
     "Fully pipelined cores, BRAM, unroll + precompute",
     "154.88 Gbps, 10.94 Mbps/slice (Kintex-7)",
     "Fully pipelined over independent inputs; no iterative-vs-unrolled study",
     "Upper bound for throughput; justifies throughput/LUT metric",
     3, 2, "Paywalled", "https://www.sciencedirect.com/science/article/abs/pii/S0141933118302758", "Partly verified (survey [17])"),
    ("F07", 2018, "Lightweight and High Performance SHA-256 using Architectural Folding and 4-2 Adder Compressor",
     "M.M. Wong, V. Pudi, A. Chattopadhyay", "IFIP/IEEE VLSI-SoC 2018, pp. 95-100", "Conference",
     "None (conf.)", "Med",
     "Area vs speed of SHA-256",
     "Folding + 4-2 compressor",
     "Best throughput/area of its time",
     "No unrolled-pair operand study",
     "Another adder-chain fix to compare against",
     3, 2, "Paywalled", "https://www.researchgate.net/publication/331796983", "Search-verified"),
    ("F08", 2002, "An FPGA based SHA-256 processor",
     "K.K. Ting, S.C.L. Yuen, K.H. Lee, P.H.W. Leong", "FPL 2002, LNCS 2438", "Conference",
     "None (CORE A)", "Med",
     "First practical FPGA SHA-256",
     "Iterative SHA-256 processor",
     "87 MB/s @88 MHz; 53 MB/s measured in-system",
     "Baseline only",
     "Our Config A follows this",
     3, 1, "Paywalled", "https://doi.org/10.1007/3-540-46117-5_60", "Verified (survey [14])"),
    ("F09", 2014, "A compact FPGA-based processor for the Secure Hash Algorithm SHA-256",
     "R. Garcia, I. Algredo-Badillo, M. Morales-Sandoval, C. Feregrino-Uribe, R. Cumplido", "Computers & Electrical Eng. 40(1)", "Journal",
     "~4.0", "Med",
     "Minimal-area SHA-256",
     "Compact processor-style SHA-256",
     "Very low slice count",
     "Area-only focus",
     "Lower bound for area",
     2, 1, "Open (author copy)", "https://ccc.inaoep.mx/~rcumplido/papers/2014-Garcia-A%20Compact%20FPGA.pdf", "Search-verified"),
    ("F10", 1991, "Retiming Synchronous Circuitry",
     "C.E. Leiserson, J.B. Saxe", "Algorithmica 6:5-35", "Journal",
     "~1.0", "High (classic)",
     "Clock period vs register placement",
     "Retiming; C-slow transformation",
     "Theory behind interleaving independent streams",
     "General theory, not SHA-256",
     "Kills the orthogonality claim (already dropped)",
     4, 4, "Paywalled", "https://doi.org/10.1007/BF01759032", "Verified (patent doc)"),
    # ---- Patents and products ----
    ("X01", 2010, "Tiny Hash Core Family for Lattice FPGA (datasheet)",
     "Helion Technology", "Commercial IP datasheet, Rev 1.0", "Product",
     "N/A", "High (commercial)",
     "Multiple authenticated streams",
     "Switches messages block-by-block on one core",
     "Sold commercially in 2010",
     "-",
     "KILLS Config C novelty",
     5, 5, "Vendor site", "https://www.heliontech.com", "Verified (patent doc)"),
    ("X02", 2025, "US20250070957A1 (IBM)",
     "IBM", "US patent application, filed 2023-08-22", "Patent",
     "N/A", "High",
     "Hash datapath throughput",
     "Background cites 'lockstep of two messages ... two-cycle pipeline'",
     "Applicant admits C/D technique is known",
     "-",
     "KILLS Config C/D novelty",
     5, 5, "Open (Google Patents)", "https://patents.google.com/patent/US20250070957A1", "Verified (patent doc)"),
    ("X03", 2025, "US12413388B2 'Methods and apparatus to hash data' (Intel)",
     "Intel", "US patent, granted 2025-09-09", "Patent",
     "N/A", "High",
     "Hash many blocks efficiently",
     "Shared message schedule; SIMD interleaved independent blocks",
     "Granted 2025",
     "-",
     "Prior art for multi-message hashing",
     4, 4, "Open (Google Patents)", "https://patents.google.com/patent/US12413388B2", "Verified (patent doc)"),
    ("X04", 2017, "Method and apparatus to process SHA-2 secure hashing algorithm (Intel SHA-NI family, e.g. US9632782)",
     "Intel", "US patents (family)", "Patent",
     "N/A", "High",
     "SHA-256 speed on CPUs",
     "SHA256RNDS2: 2 rounds per instruction with W+K PRE-ADDED",
     "Shipping in x86 CPUs",
     "CPU, not FPGA; no Fmax/threshold analysis",
     "Hoisting K+W ahead of a 2-round unit is KNOWN -> claim only the measured effect",
     5, 4, "Open (USPTO)", "https://image-ppubs.uspto.gov/dirsearch-public/print/downloadPdf/9632782", "Search-snippet (NEW)"),
]

# Open-access PDFs the download script fetches: id -> (url, filename)
OPEN_PDFS = {
    "P01": ("http://www.pertanika.upm.edu.my/resources/files/Pertanika%20PAPERS/JST%20Vol.%2030%20(1)%20Jan.%202022/32%20JST-2895-2021", "P01_Suhaili_Julai_2022.pdf"),
    "P07": ("https://www.mdpi.com/1424-8220/24/12/3908/pdf", "P07_SantosJr_2024_Sensors.pdf"),
    "P10": ("https://jisis.org/wp-content/uploads/2025/07/2025.I2.015.pdf", "P10_AlGailani_2025_JISIS.pdf"),
    "P12": ("https://www.mdpi.com/2073-431X/13/1/9/pdf", "P12_Franck_2024_Computers.pdf"),
    "P17": ("https://eprint.iacr.org/2026/1572.pdf", "P17_SHARMONY_2026.pdf"),
    "P18": ("https://eprint.iacr.org/2012/371.pdf", "P18_Gueron_Krasnov_2012.pdf"),
    "P23": ("https://eprint.iacr.org/2024/367.pdf", "P23_Saarinen_SLotH_2024.pdf"),
    "P24": ("https://eprint.iacr.org/2025/621.pdf", "P24_SPHINCSLET_2025.pdf"),
    "P25": ("https://eprint.iacr.org/2026/837.pdf", "P25_Trident_2026.pdf"),
    "P27": ("https://tches.iacr.org/index.php/TCHES/article/download/831/783/", "P27_Amiet_SPHINCS256_2020.pdf"),
    "P28": ("https://arxiv.org/pdf/2602.04415", "P28_CryptoRV_2026.pdf"),
    "P29": ("https://arxiv.org/pdf/2402.10395", "P29_Parisi_OpenTitan_2024.pdf"),
    "P30": ("https://jcoms.fesb.unist.hr/pdfs/v17n2_2021-0006_kammoun.pdf", "P30_Kammoun_2021.pdf"),
    "F05": ("https://shura.shu.ac.uk/18348/1/harris.pdf", "F05_Michail_HMAC_2012.pdf"),
    "F09": ("https://ccc.inaoep.mx/~rcumplido/papers/2014-Garcia-A%20Compact%20FPGA.pdf", "F09_Garcia_2014.pdf"),
    "X04": ("https://image-ppubs.uspto.gov/dirsearch-public/print/downloadPdf/9632782", "X04_US9632782_Intel_SHA2.pdf"),
}

# Novelty verdict per candidate claim
VERDICT = [
    ("Two-message interleaved core (Config C)", "DEAD",
     "X01 Helion 2010, X02 IBM, P17 SHARMONY 2026 (duet mode), P18 Gueron 2012, P20", 1),
    ("Unroll + pipeline / interleave combined (Config D)", "DEAD",
     "F01 McEvoy 2006, P03 Gamgam 2023, X02 IBM", 1),
    ("Operand reordering / rescheduling of T1 adds (technique)", "DEAD",
     "F02 Chaves 2006, P04 Yao 2025, X04 Intel SHA256RNDS2, F05", 1),
    ("Orthogonality of unroll and interleave depth", "DEAD (already dropped)",
     "F10 Leiserson & Saxe 1991", 1),
    ("Controlled re-test of Suhaili & Julai's Fmax rise on Xilinx (one round module, one constraint set)", "SURVIVES",
     "No paper found that isolates unroll depth as the only variable", 6),
    ("Operand order inside the unrolled pair decides pass/fail vs a pre-registered threshold (0.500 vs 0.680, R^2 = 0.998)", "SURVIVES",
     "Technique known (F02, P04, X04), but no paper measures its effect on the unrolling verdict", 6),
    ("Core vs system threshold (0.515 -> 0.624) from fixed AXI overhead", "SURVIVES (minor)",
     "Offload-overhead idea known (P29); the explicit threshold shift is ours", 4),
]
OVERALL = [
    ("Patentability", "1/10", "Every hardware technique has prior art (see Verdict rows 1-4)."),
    ("Research novelty today", "4/10", "Explanatory finding is new but rests on ngspice; Vivado numbers still pending."),
    ("Research novelty after the 4 Vivado runs", "6/10",
     "Good for a conference paper or IEEE Access-tier journal if B lands in 0.500-0.680 as predicted, or contradicts it with a reason."),
]


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Papers"
    ws.append(COLUMNS)
    for row in PAPERS:
        ws.append([("" if v is None else v) for v in row])

    head = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    wrap = Alignment(wrap_text=True, vertical="top")
    hi = PatternFill("solid", fgColor="F8CBAD")
    mid = PatternFill("solid", fgColor="FFE699")
    threat_col = COLUMNS.index("Novelty threat (1-5)") + 1
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = wrap
        t = r[threat_col - 1].value
        if t == 5:
            r[threat_col - 1].fill = hi
        elif t == 4:
            r[threat_col - 1].fill = mid
    autosize(ws, [6, 6, 42, 28, 28, 11, 16, 11, 26, 34, 32, 34, 38, 9, 9, 16, 30, 24])
    ws.freeze_panes = "D2"
    tab = Table(displayName="Papers", ref=f"A1:{get_column_letter(len(COLUMNS))}{len(PAPERS) + 1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tab)

    vs = wb.create_sheet("Novelty verdict")
    vs.append(["Candidate claim", "Status", "Killed by / evidence", "Score (1-10)"])
    for row in VERDICT:
        vs.append(list(row))
    vs.append([])
    vs.append(["Overall", "Score", "Why"])
    for row in OVERALL:
        vs.append(list(row))
    for r in vs.iter_rows():
        for c in r:
            c.alignment = wrap
    for c in vs[1]:
        c.font = Font(bold=True)
    for c in vs[len(VERDICT) + 3]:
        c.font = Font(bold=True)
    autosize(vs, [60, 22, 70, 12])

    ls = wb.create_sheet("Legend")
    for line in [
        ["Relevance (1-5)", "5 = direct prior art for Config B/C/D or the operand-scheduling finding; 1 = background"],
        ["Novelty threat (1-5)", "5 = kills a novelty claim outright; 1 = no overlap"],
        ["Impact factor", "Approximate JCR value for journals; conferences/preprints/patents have none (CORE rank or tier given). Check on JCR before quoting."],
        ["Check status", "Verified = checked against publisher record in the survey; Search-verified = title/authors/venue confirmed by web search; flags in CAPS still need checking"],
        ["PDF", "Open = free full text exists; run pdfs/download_pdfs.sh. Paywalled = fetch via VIT library / IEEE Xplore login and drop into pdfs/"],
        ["Our work", "SHA-256 on Zynq-7020: A iterative, B 2x unrolled, C 2-msg interleaved, D both; one round module, one constraint set; ngspice: delay(n) = 0.2110 + 0.3169n ns"],
    ]:
        ls.append(line)
    for r in ls.iter_rows():
        r[0].font = Font(bold=True)
        r[1].alignment = wrap
    autosize(ls, [22, 110])

    wb.save(os.path.join(HERE, "SHA256_Literature_Tracker.xlsx"))

    with open(os.path.join(HERE, "papers.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(COLUMNS)
        for row in PAPERS:
            w.writerow([("" if v is None else v) for v in row])

    with open(os.path.join(HERE, "pdfs", "download_pdfs.sh"), "w") as f:
        f.write("#!/usr/bin/env bash\n# Generated by build_tracker.py. Fetches every open-access PDF in the tracker.\n")
        f.write("set -u\ncd \"$(dirname \"$0\")\"\n")
        for pid, (url, name) in OPEN_PDFS.items():
            f.write(f"[ -s '{name}' ] || curl -fsSL -A 'Mozilla/5.0' -o '{name}' '{url}' || echo 'FAILED {pid} {url}'\n")
    os.chmod(os.path.join(HERE, "pdfs", "download_pdfs.sh"), 0o755)
    print(f"{len(PAPERS)} papers, {len(OPEN_PDFS)} open PDFs")


if __name__ == "__main__":
    main()
